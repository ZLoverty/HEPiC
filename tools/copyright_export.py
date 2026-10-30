
from pathlib import Path
from argparse import ArgumentParser
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_LINE_SPACING

EXCLUDE = {
    ".git", ".venv", "venv", "__pycache__",
    "build", "dist", ".tox", ".mypy_cache",
    ".pytest_cache", "site-packages",
}

LINES_PER_PAGE = 50
MAX_PAGES = 60


def collect_source(root):
    lines = []
    files = sorted(root.rglob("*.py"))

    for path in files:
        if any(part in EXCLUDE for part in path.relative_to(root).parts):
            continue

        content = path.read_text(encoding="utf-8-sig")
        lines.extend(content.splitlines())

    return lines


def select_source(lines):
    limit = LINES_PER_PAGE * MAX_PAGES
    if len(lines) <= limit:
        return lines
    half = limit // 2
    return lines[:half] + lines[-half:]


def generate(lines, name, version, output):
    doc = Document()
    section = doc.sections[0]

    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    header = section.header.paragraphs[0]
    header.text = f"{name} {version}"
    header.alignment = 1

    style = doc.styles["Normal"]
    style.font.name = "Consolas"
    style.font.size = Pt(8)

    fmt = style.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(12)
    fmt.keep_with_next = False
    fmt.widow_control = False

    for start in range(0, len(lines), LINES_PER_PAGE):
        if start:
            doc.add_page_break()

        chunk = lines[start:start + LINES_PER_PAGE]
        for line in chunk:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(12)
            p.add_run(line.replace("\t", "    "))

    footer = section.footer.paragraphs[0]
    footer.alignment = 1

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    doc.save(output)


def main():
    parser = ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--name", required=True)
    parser.add_argument("--version", default="V1.0")
    parser.add_argument("--output", default="source-code.docx")
    args = parser.parse_args()

    root = Path(args.source).resolve()
    if not root.is_dir():
        parser.error(f"目录不存在: {root}")

    lines = collect_source(root)
    if not lines:
        parser.error("未找到 Python 源代码")

    selected = select_source(lines)
    generate(selected, args.name, args.version, args.output)

    print(f"源代码总行数: {len(lines)}")
    print(f"导出行数: {len(selected)}")
    print(f"预计页数: {(len(selected) + 49) // 50}")
    print(f"输出文件: {args.output}")
    print("请用 Word 检查实际分页和长行折行情况。")


if __name__ == "__main__":
    main()
